"""The quote book's scorecard and per-quote economics reproduce the mock workbook's own cached results."""
import datetime as dt
import hashlib
import os
from pathlib import Path
import sys
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from compute_pricing.control.quote_book import scoring, seed

try:
    import openpyxl
except ImportError:  # the workbook oracle needs the dev dependency
    openpyxl = None

# The mock workbook is not committed. Point this at the file named in seed.SOURCE to run the workbook oracle.
WORKBOOK = Path(os.environ["QUOTE_BOOK_WORKBOOK"]) if os.environ.get("QUOTE_BOOK_WORKBOOK") else None
SKIP_REASON = "set QUOTE_BOOK_WORKBOOK to the mock workbook named in seed.SOURCE (needs openpyxl)"


def workbook_available():
    if not (openpyxl and WORKBOOK and WORKBOOK.exists()):
        return False
    if hashlib.sha256(WORKBOOK.read_bytes()).hexdigest() != seed.SOURCE_SHA256:
        raise AssertionError("QUOTE_BOOK_WORKBOOK is not the workbook the seed was generated from (SHA-256 differs)")
    return True
AS_OF = dt.date.fromisoformat(seed.AS_OF)


def close(actual, expected, places=9):
    if expected in (None, ""):
        return actual is None
    return actual is not None and abs(actual - expected) < 10 ** -places


@unittest.skipUnless(workbook_available(), SKIP_REASON)
class WorkbookOracleTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        book = openpyxl.load_workbook(WORKBOOK, data_only=True)
        cls.analysis = {row[0]: row for row in book["Quote Analysis"].iter_rows(min_row=3, max_row=40, values_only=True)}
        cls.scores = {row[1]: row for row in book["Provider Scores"].iter_rows(min_row=3, max_row=14, values_only=True)}
        cls.providers = {provider["name"]: provider for provider in seed.PROVIDERS}
        cls.grades = {name: scoring.grade_for(scoring.quality_score(provider["tiers"]))[0] for name, provider in cls.providers.items()}

    def test_provider_scores_grades_and_premiums_match_the_workbook(self):
        self.assertEqual(len(self.scores), 12)
        for name, row in self.scores.items():
            with self.subTest(provider=name):
                score = scoring.quality_score(self.providers[name]["tiers"])
                grade, premium = scoring.grade_for(score)
                self.assertTrue(close(score, row[9]), (score, row[9]))
                self.assertEqual((grade, premium), (row[10], row[11]))

    def test_every_quote_column_matches_the_workbook(self):
        self.assertEqual(len(self.analysis), 38)
        for quote in seed.QUOTES:
            row = self.analysis[quote["id"]]
            grade = self.grades[quote["provider"]]
            reference = scoring.ASSUMPTIONS["reference"].get(quote["gpu"])
            result = scoring.quote_economics(quote, grade=grade, reference=reference, as_of=AS_OF)
            with self.subTest(quote=quote["id"]):
                self.assertEqual(row[8], grade)
                self.assertEqual(result["gpus"], row[3])
                for key, column in (("contract_value", 9), ("down_payment_usd", 10), ("months_until_available", 11),
                                    ("commitment_beyond_term", 12), ("financing_cost", 13), ("scorecard_charge", 14),
                                    ("failure_probability", 15), ("exposure_charge", 17), ("risk_adjusted_scorecard", 18),
                                    ("risk_adjusted", 19), ("discount_to_reference", 21)):
                    self.assertTrue(close(result[key], row[column], 6), (key, result[key], row[column]))
                self.assertEqual(len(result["missing"]), row[22])
                self.assertEqual(" ".join(result["flags"]), row[23] or "")
                self.assertEqual("Yes" if result["eligible"] else "No", row[26])


class ScorecardTests(unittest.TestCase):
    def test_weights_form_a_whole_and_perfect_tiers_score_one_hundred(self):
        self.assertAlmostEqual(sum(weight for _, weight in scoring.CATEGORIES), 1.0)
        self.assertAlmostEqual(scoring.quality_score({name: 1 for name, _ in scoring.CATEGORIES}), 100.0)
        self.assertAlmostEqual(scoring.quality_score({name: 3 for name, _ in scoring.CATEGORIES}), 100 / 3)

    def test_grade_bands_are_inclusive_at_their_minimum(self):
        for score, expected in ((100, "A"), (85, "A"), (84.99, "B"), (70, "B"), (69.99, "C"), (55, "C"), (54.99, "D"), (0, "D")):
            self.assertEqual(scoring.grade_for(score)[0], expected, score)
        self.assertEqual([scoring.grade_for(score)[1] for score in (90, 75, 60, 10)], [0.0, 0.02, 0.05, 0.10])

    def test_ineligibility_reasons_are_stated(self):
        quote = {**seed.QUOTES[0]}  # Q001: broker, 7 of 10 key fields missing
        result = scoring.quote_economics(quote, grade="D", reference=6.65, as_of=AS_OF)
        self.assertFalse(result["eligible"])
        self.assertIn("grade D is below C", result["ineligible_reasons"])
        self.assertIn("7 of 10 key fields missing", result["ineligible_reasons"])
        clean = {**quote, **{name: "x" for name in scoring.KEY_FIELDS}, "months": 24, "down_payment": 0.25,
                 "expires": "2026-12-31"}
        self.assertTrue(scoring.quote_economics(clean, grade="A", reference=6.65, as_of=AS_OF)["eligible"])
        no_term = {**clean, "months": None}
        self.assertIn("no stated duration", scoring.quote_economics(no_term, grade="A", reference=None, as_of=AS_OF)["ineligible_reasons"])

    def test_no_reference_price_means_no_replacement_cost_and_no_discount(self):
        quote = {**seed.QUOTES[1]}  # GB300 has no public index
        result = scoring.quote_economics(quote, grade="D", reference=None, as_of=AS_OF)
        self.assertIsNone(result["discount_to_reference"])
        self.assertIsNotNone(result["risk_adjusted"])


if __name__ == "__main__":
    unittest.main()
